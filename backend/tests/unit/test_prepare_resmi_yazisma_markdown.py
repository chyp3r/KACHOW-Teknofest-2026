"""Unit tests for the Markdown-first official-correspondence preparation."""

import csv
import json
import os
import sys
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))

import prepare_resmi_yazisma_markdown as prepare  # noqa: E402


def test_front_matter_round_trip_preserves_colons_and_turkish_text():
    original = {
        "id": "BEL-001",
        "kategori": "diger_resmi_yazisma",
        "baslik": "Encümen Kararı: 2026/42",
        "rag_status": "candidate",
    }

    meta, body = prepare.split_front_matter(prepare.render_card(original, "# Başlık\n\nGövde."))

    assert meta == original
    assert body == "# Başlık\n\nGövde."


def test_semantic_anonymization_uses_context_instead_of_deleted_marker():
    text = (
        "Sayı: E-[SİLİNMİŞTİR]\n"
        "Başvuran: [SİLİNMİŞTİR]\n"
        "Telefon: 0532 123 45 67\n"
        "E-posta: ad.soyad@example.org\n"
        "[SİLİNMİŞTİR]\nDaire Başkanı"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "[SİLİNMİŞTİR]" not in anonymized
    assert "**Sayı:** [EVRAK SAYISI]" in anonymized
    assert "Başvuran: [BAŞVURU SAHİBİ]" in anonymized
    assert "[KURUM İLETİŞİM BİLGİLERİ]" in anonymized
    assert "E-posta: [E-POSTA]" in anonymized
    assert "[İMZA SAHİBİ]\nDaire Başkanı" in anonymized


def test_semantic_anonymization_repairs_legacy_placeholders_by_role():
    text = (
        "Sayın [EVRAK SAYISI]\n"
        "[EVRAK SAYISI] Milletvekili Sayın [EVRAK SAYISI]\n"
        "[EVRAK SAYISI] Milletvekili [EVRAK SAYISI]'e Ait\n"
        "Prof. Dr. [KURUM ADI]\n"
        "**VEKİLİ:** Av. [KİŞİSEL BİLGİ]\n"
        "[KİŞİSEL BİLGİ] başvuru numaralı dilekçe\n"
        "Prof. Dr. [KİŞİSEL BİLGİ]\nRektör Yardımcısı\n"
        "Katip Üye\n[KİŞİSEL BİLGİ]\n"
        "[KİŞİSEL BİLGİ]\nİstanbul Valisi"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "Sayın [KİŞİ ADI]" in anonymized
    assert "[İL ADI] Milletvekili Sayın [KİŞİ ADI]" in anonymized
    assert "[İL ADI] Milletvekili [KİŞİ ADI]'e Ait" in anonymized
    assert "Prof. Dr. [KİŞİ ADI]" in anonymized
    assert "**VEKİLİ:** Av. [VEKİL ADI]" in anonymized
    assert "[KAYIT NUMARASI] başvuru numaralı" in anonymized
    assert "Prof. Dr. [İMZA SAHİBİ]\nRektör Yardımcısı" in anonymized
    assert "Katip Üye\n[İMZA SAHİBİ]" in anonymized
    assert "[İMZA SAHİBİ]\nİstanbul Valisi" in anonymized


def test_bold_number_and_subject_on_one_line_are_preserved_as_two_fields():
    text = "**Sayı:** E-[SİLİNMİŞTİR] **Konu:** Bütçe Ödeneği Aktarımı"

    anonymized = prepare.semantic_anonymize(text)

    assert anonymized == "**Sayı:** [EVRAK SAYISI]\n**Konu:** Bütçe Ödeneği Aktarımı"


def test_public_official_document_and_decision_numbers_are_not_blindly_removed():
    text = "**Sayı:** E-12345678-2026/42\n**Esas Sayısı:** 2023/131\n**Karar Sayısı:** 2023/160"

    assert prepare.semantic_anonymize(text) == text


def test_corrupted_masked_document_number_tail_is_collapsed_without_touching_prose():
    text = "[EVRAK SAYISI]ük. 20/B-2024/13]-25650\nMetin [EVRAK SAYISI] ile ilgilidir."

    assert prepare.semantic_anonymize(text) == (
        "[EVRAK SAYISI]\nMetin [EVRAK SAYISI] ile ilgilidir."
    )


def test_semantic_anonymization_masks_generated_personal_data():
    text = "Ahmet Yılmaz (TCKN: 12345678901) adlı kişinin adresi.\nAdres: Test Mah. No: 3"

    anonymized = prepare.semantic_anonymize(text)

    assert "Ahmet Yılmaz" not in anonymized
    assert "12345678901" not in anonymized
    assert "Test Mah." not in anonymized
    assert "[KİŞİ ADI]" in anonymized
    assert "[T.C. KİMLİK NO]" in anonymized
    assert "[ADRES]" in anonymized


def test_official_contact_lines_are_generalized_without_masking_legal_prose():
    text = (
        "Bilgi için: Altunser KARAKURUMER\n"
        "Güvenevler Mah. Kuzgun Cad. No:51 Aşağı Ayrancı/ANKARA\n"
        "Santral: (0312) 466 07 37\n"
        "Mahalle muhtarlarına ödenek ödenmesi hakkında görüş istenmiştir."
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "Bilgi için: [KİŞİ ADI]" in anonymized
    assert "[KURUM ADRESİ]" in anonymized
    assert "[KURUM İLETİŞİM BİLGİLERİ]" in anonymized
    assert "Mahalle muhtarlarına ödenek" in anonymized


def test_ocr_joined_phone_fax_and_kep_footer_becomes_one_placeholder():
    text = (
        "Telefon No: 4491367 Faks No: 4491341 Burak BİLGİCİ\n"
        "KEP Adresi: kurum@example.kep.tr Telefon No: 4491367"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert anonymized == "[KURUM İLETİŞİM BİLGİLERİ]"
    assert "Burak" not in anonymized


def test_honorific_and_extended_signature_title_person_names_are_masked():
    text = (
        "İstanbul Milletvekili Sayın Fethi AÇIKEL tarafından sunulmuştur.\n\n"
        "Cevdet YILMAZ\nCumhurbaşkanı Yardımcısı"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "Sayın [KİŞİ ADI] tarafından" in anonymized
    assert "[İMZA SAHİBİ]\nCumhurbaşkanı Yardımcısı" in anonymized
    assert "Fethi" not in anonymized
    assert "Cevdet" not in anonymized


def test_role_specific_names_and_identifiers_use_semantic_placeholders():
    text = (
        "Başvuran: Ayşe Demir\n"
        "Vekili: Av. Mehmet Öztürk\n"
        "İmza Sahibi: Fatma Kaya\n"
        "Personel Sicil No: 987654\n"
        "IBAN: TR330006100519786457841326"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "Başvuran: [BAŞVURU SAHİBİ]" in anonymized
    assert "Vekili: Av. [VEKİL ADI]" in anonymized
    assert "İmza Sahibi: [İMZA SAHİBİ]" in anonymized
    assert "Personel Sicil No: [SİCİL NUMARASI]" in anonymized
    assert "IBAN: [IBAN]" in anonymized
    assert not any(name in anonymized for name in ("Ayşe", "Mehmet", "Fatma"))


def test_turkish_circumflex_vowels_are_not_truncated_mid_word():
    """``Millî``/``resmî``/``kâğıt`` use letters (â/î/û) formal Turkish
    correspondence relies on. A name/institution regex's capitalized-word
    character class silently missed them, so it stopped matching one letter
    short and left a dangling circumflex vowel glued to the untouched
    remainder (e.g. ``[KİŞİ ADI]î Eğitim Müdürü``) instead of covering the
    whole honorific+title run.
    """
    text = "Sayın İl Millî Eğitim Müdürü, başvurunuz incelenmiştir."

    anonymized = prepare.semantic_anonymize(text)

    assert anonymized == "Sayın [KİŞİ ADI], başvurunuz incelenmiştir."
    assert "î" not in anonymized.split("]")[1].split(",")[0]


def test_bench_title_is_not_masked_as_a_legal_representative():
    """``Başkanvekili`` is a court title, not an attorney."""
    text = (
        "**Başkanvekili:** Hasan Tahsin Gokcan\n"
        "**VEKİLİ:** Av. Mehmet Ozturk"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "**Başkanvekili:** [KİŞİ ADI]" in anonymized
    assert "**VEKİLİ:** Av. [VEKİL ADI]" in anonymized
    # Older runs stored the wrong placeholder; the pass must converge on rerun.
    assert prepare.semantic_anonymize("**Başkanvekili:** [VEKİL ADI]") == (
        "**Başkanvekili:** [KİŞİ ADI]"
    )
    assert prepare.semantic_anonymize(anonymized) == anonymized


def test_markdown_role_lists_ascii_names_and_context_identifiers_are_masked():
    text = (
        "**Başkan:** Kadir Ozkaya\n"
        "**Üyeler:** Basri Bagci, Kenan Yasar\n"
        "- Abone No: 397420\n"
        "- Sayaç No: 80470643\n"
        "- SGK Sicil No: 123456\n"
        "E-posta: bilgi[at]example.org\n"
        "Selin Gunes\nİmza"
    )

    anonymized = prepare.semantic_anonymize(text)

    assert "**Başkan:** [KİŞİ ADI]" in anonymized
    assert "**Üyeler:** [KİŞİ ADI]" in anonymized
    assert "- Abone No: [ABONE NUMARASI]" in anonymized
    assert "- Sayaç No: [SAYAÇ NUMARASI]" in anonymized
    assert "- SGK Sicil No: [SİCİL NUMARASI]" in anonymized
    assert "E-posta: [E-POSTA]" in anonymized
    assert "[İMZA SAHİBİ]\nİmza" in anonymized
    assert not any(
        value in anonymized
        for value in ("Kadir", "Basri", "Kenan", "397420", "80470643", "123456", "Selin")
    )


def test_contextual_audit_never_copies_sensitive_value_and_is_idempotent():
    text = "Başvuran: Ayse Demir\nAbone No: 397420\nE-posta: kisi[at]example.org"

    findings = prepare._audit_privacy_findings(text)
    anonymized = prepare.semantic_anonymize(text)

    assert findings
    assert {finding["bulgu_turu"] for finding in findings} >= {
        "rol_etiketli_kisi_adi",
        "abone_numarasi",
        "obfuscated_e_posta",
    }
    serialized = json.dumps(findings, ensure_ascii=False)
    assert "Ayse Demir" not in serialized
    assert "397420" not in serialized
    assert "kisi[at]example.org" not in serialized
    assert all(finding["satir"] >= 1 and finding["bolum"] for finding in findings)
    assert prepare.semantic_anonymize(anonymized) == anonymized


def test_unlabelled_tckn_requires_checksum_but_labelled_value_is_always_protected():
    invalid_bare = "Belge referansı 12345678901 olarak kaydedildi."
    labelled = "TCKN: 12345678901"

    assert "12345678901" in prepare.semantic_anonymize(invalid_bare)
    assert prepare.semantic_anonymize(labelled) == "TCKN: [T.C. KİMLİK NO]"
    assert prepare._valid_tckn("10000000146")
    assert prepare.semantic_anonymize("10000000146") == "[T.C. KİMLİK NO]"


def test_reported_regression_cards_have_no_automatically_fixable_privacy_finding_after_pass():
    stems = (
        "YARG-001_aym_ek_mtv_iptal_isteminin_reddi",
        "YARG-002_aym_mulkiyet_hakki_ihlali_bireysel_basvuru",
        "YARG-003_aym_makul_surede_yargilanma_hakki",
        "YARG-004_aym_iyuk_parasal_sinirlar_iptal_karari",
        "YARG-007_aym_ifade_ozgurllugu_ve_erisim_engeli",
        "bosanma_dava_anlasmali_bosanma_dava_dilekcesi_ornek_2",
        "dilekceornegi_kunye",
        "iski_abone_iptal_iskiaski_su_aboneligi_iptal_dilekcesi_ornek_2",
        "iskur_ise_kayit_iskur_issizlik_maasi_basvuru_dilekcesi_ornek_2",
        "iskur_ise_kayit_iskur_issizlik_maasi_basvuru_dilekcesi_ornek_3",
        "kdk_basvuru_kamu_denetciligi_kurumu_ombudsman_basvuru_dilekcesi_ornek_1",
        "kdk_basvuru_kamu_denetciligi_kurumu_ombudsman_basvuru_dilekcesi_ornek_2",
    )

    for stem in stems:
        paths = list(prepare.CORPUS_ROOT.rglob(f"{stem}.md"))
        assert paths, f"Regression card not found: {stem}"
        for path in paths:
            _meta, body = prepare.split_front_matter(path.read_text(encoding="utf-8"))
            sanitized = prepare.semantic_anonymize(prepare.normalize_markdown(body))
            findings = prepare._audit_privacy_findings(sanitized)
            assert not [
                finding for finding in findings if finding["otomatik_duzeltilebilir"]
            ], path.as_posix()


def test_committed_audit_manifest_never_republishes_a_raw_personal_value():
    """The shipped audit manifest must stay a structural report, not a leak."""
    manifest = prepare.ANONYMIZATION_AUDIT_MANIFEST
    if not manifest.exists():
        return

    allowed_keys = {
        "bolum",
        "bulgu_id",
        "bulgu_turu",
        "dosya",
        "duzeltme_durumu",
        "guven",
        "insan_incelemesi_gerekli",
        "kart_id",
        "maskeli_onizleme",
        "onem",
        "onerilen_yer_tutucu",
        "otomatik_duzeltilebilir",
        "satir",
    }
    for line in manifest.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        finding = json.loads(line)
        assert set(finding) <= allowed_keys, sorted(set(finding) - allowed_keys)
        assert finding["maskeli_onizleme"] == (
            f"[MASKELİ ÖNİZLEME: {finding['bulgu_turu']}]"
        )
        # ``bolum`` is the nearest heading and is the only free-text field, so
        # it is the one place a value could accidentally travel through.
        assert not prepare._EMAIL.search(finding["bolum"])
        assert not prepare._IBAN.search(finding["bolum"])
        assert not prepare._PHONE.search(finding["bolum"])
        assert not prepare._TCKN.search(finding["bolum"])


def test_parenthetical_role_qualified_label_still_masks_the_name():
    """A petition can qualify its lead label with a role synonym in
    parentheses (e.g. ``İTİRAZ EDEN (DAVACI):``). This surfaced from the
    pilot vaka generator's own output: the plain ``DAVACI:`` label was
    already masked, but the compound form left the name untouched because
    the line didn't start with a recognised label.
    """
    text = "İTİRAZ EDEN (DAVACI): Mehmet Özdemir\nADRES: Test Mah. No:5"

    anonymized = prepare.semantic_anonymize(text)

    assert "İTİRAZ EDEN (DAVACI): [KİŞİ ADI]" in anonymized
    assert "Mehmet" not in anonymized
    assert "Özdemir" not in anonymized


def test_source_institution_prefers_provenance_and_preserves_institutions(tmp_path):
    official = prepare.source_institution(
        {"kaynak": "https://www.tbmm.gov.tr/belge/1"},
        "T.C. TİCARET BAKANLIĞI",
        tmp_path / "CY-1.md",
    )
    simulation = prepare.source_institution(
        {},
        "T.C. ANKARA BÜYÜKŞEHİR BELEDİYE BAŞKANLIĞI",
        tmp_path / "ANKARA_BSB_SIMULASYON_001.md",
    )

    assert official == "Türkiye Büyük Millet Meclisi"
    assert simulation == "Ankara Büyükşehir Belediyesi"


def test_all_markdown_analysis_is_idempotent_and_accounts_for_rejected_cards(
    tmp_path, monkeypatch
):
    repo_root = tmp_path / "repo"
    corpus = repo_root / "datasets" / "resmi_yazisma"
    active = corpus / "01_ust_yazi" / "UY-1.md"
    rejected = corpus / "99_reddedilenler" / "RET-1.md"
    active.parent.mkdir(parents=True)
    rejected.parent.mkdir(parents=True)
    active.write_text(
        prepare.render_card(
            {"id": "UY-1", "kategori": "ust_yazi", "kaynak": "sentetik-sablon"},
            "# Talep\n\nBaşvuran: Ayşe Demir\nSayı: 123",
        ),
        encoding="utf-8",
    )
    rejected.write_text(
        prepare.render_card(
            {
                "id": "RET-1",
                "kategori": "diger_resmi_yazisma",
                "kaynak": "sentetik-sablon",
                "rag_status": "rejected",
            },
            "# Ret\n\nİmza Sahibi: Mehmet Öztürk",
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(prepare, "REPO_ROOT", repo_root)
    monkeypatch.setattr(prepare, "CORPUS_ROOT", corpus)

    first = prepare.anonymize_all_markdown_cards(apply=True)
    first_bytes = {path: path.read_bytes() for path in (active, rejected)}
    second = prepare.anonymize_all_markdown_cards(apply=True)

    assert len(first) == len(second) == 2
    assert {path: path.read_bytes() for path in (active, rejected)} == first_bytes
    assert "Ayşe" not in active.read_text(encoding="utf-8")
    assert "Mehmet" not in rejected.read_text(encoding="utf-8")
    assert all(record["anonimlestirme_durumu"] == "uygun" for record in second)


def test_raw_petition_uses_quarantine_copy_as_the_canonical_rag_decision(
    tmp_path, monkeypatch
):
    repo_root = tmp_path / "repo"
    corpus = repo_root / "datasets" / "resmi_yazisma"
    source = corpus / "00_gelen_kaynaklar" / "dilekce" / "article.md"
    quarantine = corpus / "99_reddedilenler" / "dilekce_makaleleri" / "article.md"
    source.parent.mkdir(parents=True)
    quarantine.parent.mkdir(parents=True)
    source.write_text(
        prepare.render_card(
            {"id": "D-1", "kategori": "dilekce", "kaynak": "https://example.test"},
            "# Açıklayıcı makale\n\nBu metin tekil dilekçe değildir.",
        ),
        encoding="utf-8",
    )
    quarantine.write_text(
        prepare.render_card(
            {
                "id": "D-1",
                "kategori": "dilekce",
                "rag_status": "reference_only",
                "ret_nedeni": "aciklayici_makale_tekil_dilekce_degil",
            },
            "# Temizlenmiş açıklama\n\nReferans metni.",
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(prepare, "REPO_ROOT", repo_root)
    monkeypatch.setattr(prepare, "CORPUS_ROOT", corpus)
    monkeypatch.setattr(prepare, "SOURCE_ROOT", corpus / "00_gelen_kaynaklar")
    monkeypatch.setattr(prepare, "REJECTED_ROOT", corpus / "99_reddedilenler")

    records = prepare.anonymize_all_markdown_cards(apply=False)
    source_record = next(record for record in records if record["path"].endswith("00_gelen_kaynaklar/dilekce/article.md"))

    assert source_record["rag_status"] == "reference_only"


def test_simulation_card_is_kept_for_tests_but_never_enters_production_rag(tmp_path):
    path = tmp_path / "ANKARA_BSB_SIMULASYON_001.md"

    status, reason = prepare._effective_rag_decision(
        path,
        {"id": "ANKARA_BSB_SIMULASYON_001", "rag_status": "candidate"},
    )

    assert status == "rejected"
    assert reason == "sentetik_simulasyon_yalniz_test"


def test_quality_gate_writes_the_simulation_verdict_into_the_card(tmp_path):
    """`curate` gates on the card's own rag_status, so the card must carry it.

    Reporting the verdict only through `_effective_rag_decision` is not
    enough: a full `--apply` rewrites the card from the raw PDF and would
    otherwise stamp a readable simulation as `candidate`, letting randomised
    synthetic letterheads into the production retrieval corpus.
    """
    body = "# Simülasyon Kararı\n\n" + ("Yeterince uzun gövde metni. " * 20)

    status, reason = prepare.assess_quality(
        body,
        source=tmp_path / "ANKARA_BSB_SIMULASYON_001.pdf",
        page_count=1,
        quality_score=0.99,
        title="Simülasyon Kararı",
    )

    assert (status, reason) == ("rejected", "sentetik_simulasyon_yalniz_test")
    assert prepare.assess_quality(
        body,
        source=tmp_path / "GERCEK_BELGE_001.pdf",
        page_count=1,
        quality_score=0.99,
        title="Gerçek Belge",
    ) == ("candidate", "")


def test_qa_sample_prefers_clean_quarantine_derivative_over_raw_petition():
    base = {
        "id": "D-1",
        "kategori": "dilekce",
        "anonimlestirme_durumu": "uygun",
    }
    raw = {
        **base,
        "path": "datasets/resmi_yazisma/00_gelen_kaynaklar/dilekce/article.md",
    }
    clean = {
        **base,
        "path": "datasets/resmi_yazisma/99_reddedilenler/dilekce_makaleleri/article.md",
    }

    selected = prepare._qa_sample([raw, clean], limit=1)

    assert selected == [clean]


def test_generated_analysis_markdown_is_not_counted_as_a_data_card(tmp_path, monkeypatch):
    (tmp_path / "RAG_VERI_ANALIZI.md").write_text("# Rapor", encoding="utf-8")
    card = tmp_path / "00_gelen_kaynaklar" / "card.md"
    card.parent.mkdir()
    card.write_text("---\nid: X-1\n---\nBelge", encoding="utf-8")
    monkeypatch.setattr(prepare, "CORPUS_ROOT", tmp_path)

    assert prepare.data_markdown_files() == [card]


def test_petition_article_cleaner_removes_site_chrome_and_comment_tail():
    text = """---
id: DILEKCE-1
---
# Abonelik İptali
Anasayfa
ANASAYFA
Abonelik İptali
Dilekceornegi
7 yıl önce
2 dakikada okunabilir
Facebook'ta Paylaş
Aboneliğin iptali için kuruma yazılı başvuru yapılır.
Bir Cevap Yaz
abonelik iptali seo etiketi
"""

    cleaned = prepare.clean_petition_article(text, "Abonelik İptali")

    assert cleaned.startswith("# Abonelik İptali")
    assert "yıl önce" not in cleaned
    assert "Facebook" not in cleaned
    assert "Bir Cevap Yaz" not in cleaned
    assert "seo etiketi" not in cleaned
    assert "yazılı başvuru" in cleaned


def test_html_extractor_uses_the_official_content_container():
    html = """
    <html><body><nav>Menü</nav><h1 class="dark">Planlı Kesinti</h1>
    <div class="text-content"><p>Sistem 12.00 ile 13.00 arasında kapalı olacaktır.</p>
    <ul><li>Kullanıcılara duyurulur.</li></ul></div><footer>Alt menü</footer></body></html>
    """.encode()

    body, title = prepare.html_to_markdown(html)

    assert title == "Planlı Kesinti"
    assert "kapalı olacaktır" in body
    assert "- Kullanıcılara duyurulur." in body
    assert "Menü" not in body
    assert "Alt menü" not in body


def test_html_extractor_ignores_visually_hidden_institution_heading():
    html = """
    <html><head><meta property="og:title" content="Gerçek Duyuru" /></head><body>
    <h1 class="visually-hidden">T.C. Bakanlık</h1>
    <div class="__header"><h2>Gerçek Duyuru</h2></div>
    <div class="__content"><p>Vatandaşlara yönelik açıklama metni.</p></div>
    </body></html>
    """.encode()

    _body, title = prepare.html_to_markdown(html)

    assert title == "Gerçek Duyuru"


def test_docx_extractor_reads_text_inside_a_floating_text_box():
    document = Document()
    paragraph = document.add_paragraph("Üst bilgi")
    text_box = OxmlElement("w:txbxContent")
    box_paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = "Metin kutusundaki resmî yazışma gövdesi."
    run.append(text)
    box_paragraph.append(run)
    text_box.append(box_paragraph)
    paragraph._p.append(text_box)
    buffer = BytesIO()
    document.save(buffer)

    body = prepare.docx_to_markdown(buffer.getvalue())

    assert "Üst bilgi" in body
    assert "Metin kutusundaki resmî yazışma gövdesi." in body


def test_infer_title_uses_a_short_first_line_for_word_templates():
    assert prepare.infer_title("Tutanak Örneği;\n\nUzun gövde.", "DY-001_tutanak") == "Tutanak Örneği"


def test_quality_gate_separates_atomic_examples_from_reference_documents(tmp_path):
    body = "Resmî yazışma gövdesi. " * 20

    assert prepare.assess_quality(body, source=tmp_path / "ornek.pdf") == ("candidate", "")
    assert prepare.assess_quality("çok kısa", source=tmp_path / "ornek.pdf") == (
        "rejected",
        "yetersiz_metin",
    )
    assert prepare.assess_quality(body, source=tmp_path / "ornek.pdf", page_count=32) == (
        "reference_only",
        "tekil_yazisma_ornegi_degil",
    )


def test_normalizer_repairs_only_question_marks_used_as_apostrophes():
    text = "Türkiye?nin açıklaması nedir? Doğu Anadolu?da uygulanır."

    assert prepare.normalize_markdown(text) == (
        "Türkiye'nin açıklaması nedir? Doğu Anadolu'da uygulanır."
    )


def test_quality_gate_rejects_mojibake_and_corrupted_titles(tmp_path):
    body = "Resmî yazışma gövdesi. " * 20

    assert prepare.assess_quality(
        body + "\nDoğrulama Adresi: Âhüps/example",
        source=tmp_path / "ornek.pdf",
        title="Geçerli başlık",
    ) == ("rejected", "ocr_karakter_bozulmasi")
    assert prepare.assess_quality(
        body,
        source=tmp_path / "ornek.pdf",
        title=", Â Sayı: E-12007 03-1$23865",
    ) == ("rejected", "ocr_karakter_bozulmasi")


def test_front_matter_completion_adds_every_required_field(tmp_path, monkeypatch):
    repo_root = tmp_path / "repo"
    card = repo_root / "datasets" / "resmi_yazisma" / "01_ust_yazi" / "05_teklif" / "UY-1.md"
    monkeypatch.setattr(prepare, "REPO_ROOT", repo_root)

    completed = prepare.complete_front_matter(
        {
            "id": "UY-1",
            "kategori": "ust_yazi",
            "niyet": "05_teklif",
        },
        "# Proje Teklifi\n\nGövde.",
        card,
    )

    assert {
        "id",
        "kategori",
        "alt_kategori",
        "baslik",
        "kaynak",
        "dogrulama",
    } <= completed.keys()
    assert completed["alt_kategori"] == "05_teklif"
    assert completed["baslik"] == "Proje Teklifi"


_OS_MECLIS_BODY = (
    "Söz konusu meclis kararı, 5393 sayılı Belediye Kanununun ilgili maddeleri "
    "uyarınca oy birliği ile kabul edilmiştir."
)
_OS_HIFZISSIHHA_BODY = (
    "İl Umumi Hıfzıssıhha Kurulu, Vali başkanlığında toplanarak aşağıdaki "
    "kararları almıştır: İl genelinde halk sağlığını tehdit eden unsurlara "
    "karşı denetimler artırılacaktır."
)
_OS_UYGUNLUK_BODY = (
    "İlgi kayıtlı yazınız incelenmiş olup, talep edilen hususlar mevzuat "
    "çerçevesinde değerlendirilmiştir. Kurumumuzca yapılan inceleme neticesinde "
    "belirtilen işlemlerin uygun olduğu mütalaa edilmiştir."
)
_OS_BILGI_EDINME_BODY = (
    "Başvurunuz, 4982 sayılı Bilgi Edinme Hakkı Kanunu kapsamında incelenmiştir."
)


def test_os_coherence_gate_rejects_the_reported_budget_health_mismatch():
    assert prepare.os_is_coherent(
        "Bütçe Ödeneği Aktarımı Hakkında",
        _OS_HIFZISSIHHA_BODY,
        kategori="03_bilgilendirme_metni",
        kurum="T.C. İzmir Valiliği",
    ) == (False, "hifzissihha_karari", "baslik_govde_uyumsuzlugu")

    assert prepare.os_is_coherent(
        "Halk Sağlığı Tedbirleri Hakkında",
        _OS_HIFZISSIHHA_BODY,
        kategori="03_bilgilendirme_metni",
        kurum="T.C. İzmir Valiliği",
    ) == (True, "hifzissihha_karari", "")


def test_os_coherence_gate_is_fail_closed_for_unknown_bodies():
    """The old gate funnelled unrecognised bodies into a "genel" bucket that
    returned True unconditionally, so any body pattern the matcher did not
    literally know about was waved straight through to ``candidate``."""
    assert prepare.os_body_kind("Tamamen tanınmayan bir gövde metni.") == "taninmayan"
    assert prepare.os_is_coherent(
        "İmar Planı Değişikliği Hakkında",
        "Tamamen tanınmayan bir gövde metni.",
        kategori="04_diger_resmi_yazisma",
        kurum="T.C. Ankara Büyükşehir Belediye Başkanlığı",
    ) == (False, "taninmayan", "taninmayan_govde_kalibi")


def test_os_gate_rejects_each_reported_mismatch_class():
    """The three cards the data owner flagged, each failing a different one
    of the generator's three independent random draws."""
    # OS-02-011: a municipal council vote filed under a Sayıştay-report title.
    assert prepare.os_is_coherent(
        "Sayıştay Denetim Raporu Hakkında",
        _OS_MECLIS_BODY,
        kategori="02_cevap_yazisi itiraz_cevabi",
        kurum="T.C. Karşıyaka Kaymakamlığı",
    ) == (False, "belediye_meclis_karari", "baslik_govde_uyumsuzlugu")

    # OS-01-010: a generic "found appropriate" opinion answering an objection.
    assert prepare.os_is_coherent(
        "Kamu İhale Kurumu İtirazı Hakkında",
        _OS_UYGUNLUK_BODY,
        kategori="01_ust_yazi ust_yazi",
        kurum="T.C. Millî Eğitim Bakanlığı",
    ) == (False, "genel_uygunluk_gorusu", "baslik_govde_uyumsuzlugu")

    # OS-04-032: title and body agree, but a bilgi-edinme *reply* is filed
    # under "diğer resmî yazışma" rather than "cevap yazısı".
    assert prepare.os_is_coherent(
        "Bilgi Edinme Başvurusu Cevabı Hakkında",
        _OS_BILGI_EDINME_BODY,
        kategori="04_diger_resmi_yazisma diger_resmi_yazisma",
        kurum="T.C. Sosyal Güvenlik Kurumu Başkanlığı",
    ) == (False, "bilgi_edinme_cevabi", "kategori_govde_uyumsuzlugu")


def test_os_gate_rejects_a_body_whose_deciding_organ_contradicts_the_letterhead():
    """Only a municipality passes a "belediye meclis kararı" -- the title and
    category can both line up and the card still be nonsense."""
    assert prepare.os_is_coherent(
        "İmar Planı Değişikliği Hakkında",
        _OS_MECLIS_BODY,
        kategori="04_diger_resmi_yazisma meclis_karari",
        kurum="T.C. Anayasa Mahkemesi",
    ) == (False, "belediye_meclis_karari", "kurum_govde_uyumsuzlugu")

    assert prepare.os_is_coherent(
        "İmar Planı Değişikliği Hakkında",
        _OS_MECLIS_BODY,
        kategori="04_diger_resmi_yazisma meclis_karari",
        kurum="T.C. Ankara Büyükşehir Belediye Başkanlığı",
    ) == (True, "belediye_meclis_karari", "")

    # OS-01-133: the body says "Bakanlığımızca yürütülen projeler", which a
    # constitutional court never writes about itself.
    proje_body = (
        "Bakanlığımızca yürütülen projeler kapsamında, ekte sunulan raporların "
        "ivedilikle incelenerek sonucundan tarafımıza bilgi verilmesi hususunda "
        "gereğini rica ederim."
    )
    assert prepare.os_is_coherent(
        "Kentsel Dönüşüm Projesi Hakkında",
        proje_body,
        kategori="01_ust_yazi ust_yazi",
        kurum="T.C. Anayasa Mahkemesi",
    ) == (False, "proje_rapor_iletimi", "kurum_govde_uyumsuzlugu")


def test_os_gate_matches_titles_with_turkish_dotted_capitals():
    """``"İ".casefold()`` expands to ``i`` plus a combining dot, so a plain
    casefold comparison silently misses every title starting with İ."""
    assert prepare._fold_tr("İmar Planı") == "imar planı"
    assert prepare.os_is_coherent(
        "İmar Planı Değişikliği Hakkında",
        _OS_UYGUNLUK_BODY,
        kategori="02_cevap_yazisi cevap_yazisi",
        kurum="T.C. İzmir Valiliği",
    ) == (True, "genel_uygunluk_gorusu", "")


def test_existing_same_stem_card_is_preferred_over_creating_a_second_card(tmp_path):
    source = tmp_path / "00_gelen_kaynaklar" / "BM-001_duyuru.html"
    companion = tmp_path / "03_bilgilendirme_metni" / "03_uyari" / "BM-001_duyuru.md"
    index = {source.stem.casefold(): [companion]}

    assert prepare.target_for_source(source, index) == companion


def test_source_without_catalog_card_gets_an_adjacent_markdown_file(tmp_path):
    source = tmp_path / "pdf" / "MEB_SIMULASYON_001.pdf"

    assert prepare.target_for_source(source, {}) == source.with_suffix(".md")


def test_source_institution_prefers_curated_metadata_over_ocr_unit_heading(tmp_path):
    card = tmp_path / "CY-001.md"
    body = "T.C.\nHAZİNE VE MALİYE BAKANLIĞI\nStrateji Geliştirme Başkanlığı"

    institution = prepare.source_institution(
        {"kurum": "Türkiye Büyük Millet Meclisinde yayımlanan kurum cevabı"},
        body,
        card,
    )

    assert institution == "Türkiye Büyük Millet Meclisi"


def test_analysis_outputs_write_statistics_manifest_and_balanced_qa(tmp_path, monkeypatch):
    corpus_root = tmp_path / "datasets" / "resmi_yazisma"
    monkeypatch.setattr(prepare, "ANONYMIZATION_MANIFEST", corpus_root / "manifest.jsonl")
    monkeypatch.setattr(
        prepare,
        "ANONYMIZATION_AUDIT_MANIFEST",
        corpus_root / "audit-manifest.jsonl",
    )
    monkeypatch.setattr(prepare, "STATISTICS_JSON", corpus_root / "statistics.json")
    monkeypatch.setattr(prepare, "STATISTICS_MD", corpus_root / "statistics.md")
    monkeypatch.setattr(prepare, "QA_MANIFEST", corpus_root / "qa.csv")
    monkeypatch.setattr(
        prepare,
        "_source_files",
        lambda: [corpus_root / "source.pdf", corpus_root / "source.html"],
    )
    corpus_root.mkdir(parents=True)
    records = []
    for index in range(105):
        category = "ust_yazi" if index % 2 == 0 else "cevap_yazisi"
        records.append(
            {
                "path": f"datasets/resmi_yazisma/01_ust_yazi/UY-{index:03}.md",
                "id": f"UY-{index:03}",
                "kategori": category,
                "kaynak_kurum": "Örnek Kurum",
                "kaynak_anahtari": f"source-{index}",
                "rag_status": "candidate",
                "anonimlestirme_durumu": "uygun",
                "neden": "",
                "anonimlestirilmis": True,
                "anonimlestirilen_alanlar": {"KİŞİ ADI": 1},
                "kalan_pii_turleri": [],
                "denetim_bulgulari": [],
                "kalan_baglamsal_bulgu_turleri": [],
            }
        )

    statistics = prepare.write_analysis_outputs(records, apply=True)

    assert statistics["markdown_kaydi"] == 105
    assert statistics["aktif_korpus_kaydi"] == 105
    assert statistics["karantina_kaydi"] == 0
    assert statistics["ham_kaynak_turu_dagilimi"] == {"html": 1, "pdf": 1}
    assert json.loads(prepare.STATISTICS_JSON.read_text(encoding="utf-8"))["tekil_belge"] == 105
    assert len(prepare.ANONYMIZATION_MANIFEST.read_text(encoding="utf-8").splitlines()) == 105
    assert prepare.ANONYMIZATION_AUDIT_MANIFEST.read_text(encoding="utf-8") == ""
    with prepare.QA_MANIFEST.open(encoding="utf-8", newline="") as handle:
        qa_rows = list(csv.DictReader(handle))
    assert len(qa_rows) == 100
    assert {row["kategori"] for row in qa_rows} == {"ust_yazi", "cevap_yazisi"}
